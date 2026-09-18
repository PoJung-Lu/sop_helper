from __future__ import annotations

import json
import sys
import tempfile
import uuid
from dataclasses import dataclass, asdict
from pathlib import Path

from PySide6.QtCore import QPointF, QRectF, Qt
from PySide6.QtGui import QColor, QPainter, QPen, QPixmap, QFont
from PySide6.QtWidgets import (
    QApplication, QComboBox, QDialog, QDialogButtonBox, QFileDialog,
    QFormLayout, QGroupBox, QLabel, QListWidget, QListWidgetItem,
    QMainWindow, QMessageBox, QPushButton, QSpinBox, QDoubleSpinBox,
    QSplitter, QVBoxLayout, QHBoxLayout, QGraphicsView, QGraphicsScene,
    QCheckBox, QStackedWidget
)
from openpyxl import load_workbook
import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

try:
    import win32clipboard
    import win32con
except ImportError:
    win32clipboard = None
    win32con = None

try:
    import win32com.client
except ImportError:
    win32com = None

APP_TITLE = "SOP_Helper MVP v0.9.7"
TORQUE_UNITS = ["kgf-cm", "N-M", "無需扭力", "鎖緊就好"]
NO_TORQUE = "無需扭力"
LOCK_ONLY = "鎖緊就好"
TEXT_COLOR = (0, 0, 0, 255)
TORQUE_BOX_COLOR = (210, 0, 0, 255)
MARKER_COLOR = (220, 0, 0, 255)
WINDOWS_FONT_FILE = r"C:\Windows\Fonts\msjh.ttc"
WINDOWS_FONT_FAMILY = "Microsoft JhengHei"
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp"}
PPTX_MARKER_RADIUS_RATIO = 0.10 / 8.5


@dataclass
class Component:
    part_name: str
    part_number: str
    size: str
    required_torque: float | None
    torque_unit: str
    applicable_model: str = "Universal"
    reference_image: str = ""
    requires_oring: bool = False
    oring_spec: str = ""
    requires_silicone_grease: bool = False
    silicone_grease_note: str = ""
    torque_min: float | None = None
    torque_max: float | None = None
    notes: str = ""

    @property
    def has_torque_spec(self):
        return self.required_torque is not None and bool(self.torque_unit)


@dataclass
class Annotation:
    annotation_id: str
    component_index: int
    x: float
    y: float
    sequence: int
    torque: float | None
    torque_unit: str
    marker_only: bool = False
    label_x: float = 0.0
    label_y: float = 0.0


class AnnotationDialog(QDialog):
    def __init__(self, component, default_sequence, used_sequences, parent=None):
        super().__init__(parent)
        self.setWindowTitle("設定標記")
        self.component = component
        self.used_sequences = set(used_sequences)
        form = QFormLayout(self)
        form.addRow("料件", QLabel(f"{component.part_name} ({component.part_number})\n尺寸：{component.size}"))

        self.marker_only = QCheckBox("僅新增標記（輸出不顯示名稱、扭力或其他細節）")
        form.addRow(self.marker_only)

        self.torque = QDoubleSpinBox()
        self.torque.setRange(0, 1_000_000)
        self.torque.setDecimals(4)
        if component.required_torque is not None:
            self.torque.setValue(component.required_torque)
        form.addRow("本次扭力", self.torque)

        self.unit = QComboBox()
        self.unit.addItems(TORQUE_UNITS)
        if component.has_torque_spec and component.torque_unit in TORQUE_UNITS:
            self.unit.setCurrentText(component.torque_unit)
        elif not component.has_torque_spec:
            self.unit.setCurrentText(NO_TORQUE)
            self.torque.setEnabled(False)
        form.addRow("單位", self.unit)
        self.unit.currentTextChanged.connect(self.on_unit_changed)
        self.torque.valueChanged.connect(self.on_torque_changed)

        self.sequence = QSpinBox()
        self.sequence.setRange(1, 999)
        self.sequence.setValue(default_sequence)
        form.addRow("標記順序", self.sequence)

        hint = f"{component.required_torque:g} {component.torque_unit}" if component.has_torque_spec else NO_TORQUE
        form.addRow("規格提示", QLabel(
            f"規格扭力：{hint}\n機型：{component.applicable_model}\n"
            f"O-ring：{component.oring_spec if component.requires_oring else '無'}\n"
            f"塗矽油：{component.silicone_grease_note if component.requires_silicone_grease else '無'}"
        ))

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.validate_and_accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def on_unit_changed(self, unit):
        numeric = unit in {"kgf-cm", "N-M"}
        self.torque.setEnabled(numeric)
        if unit in {NO_TORQUE, LOCK_ONLY}:
            self.torque.setValue(0)

    def on_torque_changed(self, value):
        if value == 0 and self.unit.currentText() in {"kgf-cm", "N-M"}:
            self.unit.blockSignals(True)
            self.unit.setCurrentText(NO_TORQUE)
            self.unit.blockSignals(False)
            self.torque.setEnabled(False)

    def validate_and_accept(self):
        sequence = self.sequence.value()
        if sequence in self.used_sequences:
            QMessageBox.warning(self, "標記順序重複", f"標記順序 {sequence} 已存在，請改用其他數字。")
            return
        self.accept()

    def values(self):
        unit = self.unit.currentText()
        if unit == NO_TORQUE or self.torque.value() == 0:
            return None, NO_TORQUE, self.sequence.value(), self.marker_only.isChecked()
        return self.torque.value(), unit, self.sequence.value(), self.marker_only.isChecked()


class AnnotationCanvas(QGraphicsView):
    def __init__(self, owner):
        super().__init__()
        self.owner = owner
        self.scene = QGraphicsScene(self)
        self.setScene(self.scene)
        self.setRenderHints(QPainter.RenderHint.Antialiasing)
        self.setDragMode(QGraphicsView.DragMode.NoDrag)
        self.setBackgroundBrush(QColor("#eeeeee"))
        self.image_item = None
        self.zoom_factor = 1.0
        self._panning = False
        self._pan_start = None
        self._dragging_annotation_id = None
        self.setAcceptDrops(True)

    def set_image(self, path):
        self.scene.clear()
        self.owner.annotation_items.clear()
        self.image_item = None
        pixmap = QPixmap(path)
        if pixmap.isNull():
            raise ValueError("無法讀取圖片")
        self.image_item = self.scene.addPixmap(pixmap)
        self.scene.setSceneRect(QRectF(pixmap.rect()))
        self.fit_to_window()

    def fit_to_window(self):
        if self.image_item:
            self.resetTransform()
            self.fitInView(self.scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)
            self.zoom_factor = 1.0
            self.owner.update_zoom_label(self.zoom_factor)

    def set_zoom_at(self, factor, anchor_view_pos=None):
        factor = max(0.25, min(4.0, factor))
        if anchor_view_pos is None:
            anchor_scene = self.mapToScene(self.viewport().rect().center())
        else:
            anchor_scene = self.mapToScene(anchor_view_pos)
        self.resetTransform()
        self.scale(factor, factor)
        self.zoom_factor = factor
        new_view_pos = self.mapFromScene(anchor_scene)
        delta = new_view_pos - (anchor_view_pos if anchor_view_pos else self.viewport().rect().center())
        self.horizontalScrollBar().setValue(self.horizontalScrollBar().value() + delta.x())
        self.verticalScrollBar().setValue(self.verticalScrollBar().value() + delta.y())
        self.owner.update_zoom_label(factor)

    def wheelEvent(self, event):
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            step = 1.15 if event.angleDelta().y() > 0 else 1 / 1.15
            self.set_zoom_at(self.zoom_factor * step, event.position().toPoint())
        else:
            super().wheelEvent(event)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.image_item and self.zoom_factor == 1.0:
            self.fit_to_window()

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.MiddleButton:
            self._panning = True
            self._pan_start = event.position().toPoint()
            self.setCursor(Qt.CursorShape.ClosedHandCursor)
            return
        if event.button() == Qt.MouseButton.LeftButton and self.image_item:
            pos = self.mapToScene(event.position().toPoint())
            if self.scene.sceneRect().contains(pos):
                if self.owner.has_component_selected():
                    self.owner.add_annotation_at(pos)
                else:
                    hit = self.owner.find_marker_at(pos)
                    if hit is not None:
                        self._dragging_annotation_id = hit.annotation_id
                        self.owner.select_marker(hit.annotation_id)
                    else:
                        self.owner.try_select_marker_at(pos)
                return
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if self._panning and self._pan_start is not None:
            delta = event.position().toPoint() - self._pan_start
            self._pan_start = event.position().toPoint()
            self.horizontalScrollBar().setValue(self.horizontalScrollBar().value() - delta.x())
            self.verticalScrollBar().setValue(self.verticalScrollBar().value() - delta.y())
            return
        if self._dragging_annotation_id is not None:
            pos = self.mapToScene(event.position().toPoint())
            rect = self.scene.sceneRect()
            if rect.contains(pos):
                self.owner.update_annotation_position(self._dragging_annotation_id, pos.x()/rect.width(), pos.y()/rect.height())
            return
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.MiddleButton:
            self._panning = False
            self._pan_start = None
            self.setCursor(Qt.CursorShape.ArrowCursor)
            return
        if event.button() == Qt.MouseButton.LeftButton and self._dragging_annotation_id is not None:
            self._dragging_annotation_id = None
            return
        super().mouseReleaseEvent(event)

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self.image_item and not self.owner.has_component_selected():
            pos = self.mapToScene(event.position().toPoint())
            hit = self.owner.find_marker_at(pos)
            if hit is not None:
                self.owner.edit_annotation(hit.annotation_id)
                return
        super().mouseDoubleClickEvent(event)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            for url in event.mimeData().urls():
                if Path(url.toLocalFile()).suffix.lower() in IMAGE_EXTENSIONS:
                    event.acceptProposedAction()
                    return
        event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        urls = event.mimeData().urls()
        image_path = None
        for url in urls:
            local_path = url.toLocalFile()
            if Path(local_path).suffix.lower() in IMAGE_EXTENSIONS:
                image_path = local_path
                break
        if image_path:
            self.owner.load_image_from_path(image_path)
            event.acceptProposedAction()
        else:
            QMessageBox.information(self, "不支援的檔案", "請拖曳 PNG、JPG 或 BMP 圖片檔案。")
            event.ignore()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(1550, 900)
        self.components = []
        self.annotations = []
        self.annotation_items = []
        self.selected_annotation_id = None
        self.image_path = ""
        # 明確自行維護「是否已選取料件」狀態，不依賴 QListWidget 的 isSelected() 時序。
        self._component_selected = False
        self._last_pressed_component_row = -1
        self.canvas = AnnotationCanvas(self)
        self.component_list = QListWidget()
        self.component_list.itemDoubleClicked.connect(self.add_selected_component)
        self.component_list.itemPressed.connect(self.on_component_item_pressed)
        self.status = QLabel("已載入。未選取料件時可自由瀏覽圖片，選取料件後點擊圖片可新增標記。")
        self.mode_hint = QLabel("目前模式：瀏覽 / 選取")
        self.mode_hint.setWordWrap(True)
        self.zoom_label = QLabel("縮放：100%")
        self.marker_list = QListWidget()
        self.marker_list.itemClicked.connect(self.on_marker_list_clicked)
        self.preview_label = QLabel("尚未產生輸出預覽")
        self.preview_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_label.setStyleSheet("background:#eeeeee; color:#555;")
        self.right_stack = QStackedWidget()
        self.right_stack.addWidget(self.marker_list)
        self.right_stack.addWidget(self.preview_label)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(self.build_left_panel())
        splitter.addWidget(self.build_center_panel())
        splitter.addWidget(self.build_right_panel())
        splitter.setSizes([300, 850, 400])
        self.setCentralWidget(splitter)
        self.statusBar().addWidget(self.status)

    def build_left_panel(self):
        box = QGroupBox("資料與料件")
        layout = QVBoxLayout(box)
        for text, slot in [("載入主圖片", self.load_image), ("載入 Excel / YAML", self.load_spec)]:
            b = QPushButton(text); b.clicked.connect(slot); layout.addWidget(b)
        component_hint_label = QLabel("可用料件（點擊選取後點圖新增標記；再點一次取消選取）")
        component_hint_label.setWordWrap(True)
        layout.addWidget(component_hint_label)
        layout.addWidget(self.component_list, 1)
        b = QPushButton("使用選取料件新增標記"); b.clicked.connect(self.add_selected_component); layout.addWidget(b)
        b = QPushButton("取消選取料件（切換為瀏覽模式）"); b.clicked.connect(self.clear_component_selection); layout.addWidget(b)
        b = QPushButton("刪除選取的標記"); b.clicked.connect(self.delete_selected_annotation); layout.addWidget(b)
        layout.addWidget(self.mode_hint)
        return box

    def build_center_panel(self):
        box = QGroupBox("主圖與標記")
        layout = QVBoxLayout(box)
        toolbar = QHBoxLayout()
        for text, slot in [("適合視窗", self.fit_view), ("−", self.zoom_out), ("100%", self.zoom_100), ("＋", self.zoom_in)]:
            b = QPushButton(text); b.clicked.connect(slot); toolbar.addWidget(b)
        toolbar.addWidget(self.zoom_label)
        toolbar.addStretch()
        layout.addLayout(toolbar)
        layout.addWidget(self.canvas, 1)
        hint_label = QLabel("操作提示：可直接把圖片檔案拖曳到畫布載入；Ctrl+滾輪＝以滑鼠位置縮放；中鍵拖曳＝平移；瀏覽模式下拖曳既有標記可移動位置，雙擊可編輯內容")
        hint_label.setWordWrap(True)
        layout.addWidget(hint_label)
        return box

    def build_right_panel(self):
        box = QGroupBox("標記清單 / 輸出預覽")
        layout = QVBoxLayout(box)
        self.toggle_preview_btn = QPushButton("顯示輸出預覽")
        self.toggle_preview_btn.clicked.connect(self.toggle_preview)
        layout.addWidget(self.toggle_preview_btn)
        layout.addWidget(self.right_stack, 1)
        actions = [
            ("複製為 PowerPoint 物件（可編輯）", self.copy_pptx_shapes),
            ("複製合成圖片（僅供備援）", self.copy_composite_image),
            ("另存 PNG", self.save_png),
            ("另存可編輯 PPTX", self.save_pptx),
            ("另存專案 JSON", self.save_project),
        ]
        for text, slot in actions:
            b = QPushButton(text); b.clicked.connect(slot); layout.addWidget(b)
        return box

    def toggle_preview(self):
        if self.right_stack.currentIndex() == 0:
            self.update_preview()
            self.right_stack.setCurrentIndex(1)
            self.toggle_preview_btn.setText("顯示標記清單")
        else:
            self.right_stack.setCurrentIndex(0)
            self.toggle_preview_btn.setText("顯示輸出預覽")

    def on_component_item_pressed(self, item):
        row = self.component_list.row(item)
        if row == self._last_pressed_component_row and self._component_selected:
            # 再次點擊目前已選取的同一項 -> 取消選取，回到瀏覽/選取模式。
            self.clear_component_selection()
        else:
            self._last_pressed_component_row = row
            self._component_selected = True
            self.update_mode_hint()

    def clear_component_selection(self):
        self.component_list.clearSelection()
        self.component_list.setCurrentItem(None)
        self._last_pressed_component_row = -1
        self._component_selected = False
        self.update_mode_hint()

    def update_mode_hint(self):
        if self.has_component_selected():
            self.mode_hint.setText("目前模式：新增標記（點擊圖片建立標記）")
        else:
            self.mode_hint.setText("目前模式：瀏覽 / 選取（點擊選取、拖曳移動、雙擊編輯既有標記；中鍵拖曳可平移畫布）")

    def has_component_selected(self):
        return self._component_selected and self.component_list.currentItem() is not None

    def update_zoom_label(self, factor):
        self.zoom_label.setText(f"縮放：{factor * 100:.0f}%")

    def fit_view(self): self.canvas.fit_to_window()
    def zoom_out(self): self.canvas.set_zoom_at(self.canvas.zoom_factor * 0.8)
    def zoom_100(self): self.canvas.set_zoom_at(1.0)
    def zoom_in(self): self.canvas.set_zoom_at(self.canvas.zoom_factor * 1.25)

    def load_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "選擇主圖片", "", "Images (*.png *.jpg *.jpeg *.bmp)")
        if not path: return
        self.load_image_from_path(path)

    def load_image_from_path(self, path):
        if self.annotations:
            reply = QMessageBox.question(
                self, "確認替換圖片",
                "載入新圖片會清除目前所有標記，是否繼續？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No,
            )
            if reply != QMessageBox.StandardButton.Yes:
                return
        try:
            self.canvas.set_image(path)
            self.image_path = path
            self.annotations = []
            self.selected_annotation_id = None
            self.rebuild_marker_list()
            self.status.setText(f"已載入圖片：{Path(path).name}")
        except Exception as exc:
            QMessageBox.critical(self, "載入失敗", str(exc))

    def load_spec(self):
        path, _ = QFileDialog.getOpenFileName(self, "選擇規格檔", "", "Excel/YAML (*.xlsx *.yaml *.yml)")
        if not path: return
        try:
            self.components, warnings = load_components(path)
            self.component_list.clear()
            self.clear_component_selection()
            for index, c in enumerate(self.components):
                torque = f"{c.required_torque:g} {c.torque_unit}" if c.has_torque_spec else NO_TORQUE
                item = QListWidgetItem(f"{c.part_name} | {c.part_number} | {torque}")
                item.setData(Qt.ItemDataRole.UserRole, index); self.component_list.addItem(item)
            if warnings: QMessageBox.information(self, "載入完成，附帶提示", "\n".join(warnings))
            self.status.setText(f"已載入 {len(self.components)} 筆料件規格")
        except Exception as exc: QMessageBox.critical(self, "規格載入失敗", str(exc))

    def selected_component_index(self):
        if not self.has_component_selected():
            return None
        item = self.component_list.currentItem()
        return item.data(Qt.ItemDataRole.UserRole) if item else None

    def next_sequence(self): return max((a.sequence for a in self.annotations), default=0) + 1

    def add_selected_component(self):
        index = self.selected_component_index()
        if index is None: QMessageBox.information(self, "提示", "請先在左側選擇料件"); return
        if not self.canvas.image_item: QMessageBox.information(self, "提示", "請先載入主圖片"); return
        self.add_annotation_at(self.canvas.scene.sceneRect().center(), index)

    def add_annotation_at(self, pos, component_index=None):
        if component_index is None: component_index = self.selected_component_index()
        if component_index is None: return
        dialog = AnnotationDialog(self.components[component_index], self.next_sequence(), [a.sequence for a in self.annotations], self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        torque, unit, sequence, marker_only = dialog.values()
        rect = self.canvas.scene.sceneRect()
        annotation = Annotation(f"F{len(self.annotations)+1:03d}", component_index, pos.x()/rect.width(), pos.y()/rect.height(), sequence, torque, unit, marker_only)
        self.annotations.append(annotation)
        self.rebuild_marker_list()
        self.update_canvas_marks()

    def find_marker_at(self, scene_pos, tolerance=22):
        rect = self.canvas.scene.sceneRect()
        best, best_dist = None, tolerance
        for a in self.annotations:
            x, y = a.x * rect.width(), a.y * rect.height()
            dist = ((scene_pos.x() - x) ** 2 + (scene_pos.y() - y) ** 2) ** 0.5
            if dist <= best_dist:
                best, best_dist = a, dist
        return best

    def try_select_marker_at(self, scene_pos, tolerance=22):
        best = self.find_marker_at(scene_pos, tolerance)
        self.selected_annotation_id = best.annotation_id if best else None
        self.update_canvas_marks()
        if best:
            self.status.setText(f"已選取標記 {best.sequence}（{self.components[best.component_index].part_name}），可拖曳移動或雙擊編輯")
            self.select_marker_in_list(best.annotation_id)
        else:
            self.status.setText("未選取任何標記")

    def select_marker(self, annotation_id):
        self.selected_annotation_id = annotation_id
        self.update_canvas_marks()
        self.select_marker_in_list(annotation_id)

    def find_annotation_by_id(self, annotation_id):
        for a in self.annotations:
            if a.annotation_id == annotation_id:
                return a
        return None

    def update_annotation_position(self, annotation_id, x, y):
        a = self.find_annotation_by_id(annotation_id)
        if a is None: return
        a.x, a.y = max(0.0, min(1.0, x)), max(0.0, min(1.0, y))
        self.update_canvas_marks()

    def edit_annotation(self, annotation_id):
        a = self.find_annotation_by_id(annotation_id)
        if a is None: return
        component = self.components[a.component_index]
        used = [s.sequence for s in self.annotations if s.annotation_id != annotation_id]
        dialog = AnnotationDialog(component, a.sequence, used, self)
        dialog.marker_only.setChecked(a.marker_only)
        if a.torque is not None:
            dialog.torque.setValue(a.torque)
        dialog.unit.setCurrentText(a.torque_unit)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        torque, unit, sequence, marker_only = dialog.values()
        a.torque, a.torque_unit, a.sequence, a.marker_only = torque, unit, sequence, marker_only
        self.rebuild_marker_list()
        self.update_canvas_marks()
        self.status.setText(f"已更新標記 {a.sequence}（{component.part_name}）")

    def delete_selected_annotation(self):
        if self.selected_annotation_id is None:
            QMessageBox.information(self, "提示", "請先在畫布或清單選取一個標記")
            return
        a = self.find_annotation_by_id(self.selected_annotation_id)
        if a is None: return
        self.annotations.remove(a)
        self.selected_annotation_id = None
        self.rebuild_marker_list()
        self.update_canvas_marks()
        self.status.setText(f"已刪除標記 {a.sequence}")

    def select_marker_in_list(self, annotation_id):
        for i in range(self.marker_list.count()):
            item = self.marker_list.item(i)
            if item.data(Qt.ItemDataRole.UserRole) == annotation_id:
                self.marker_list.setCurrentItem(item)
                break

    def on_marker_list_clicked(self, item):
        annotation_id = item.data(Qt.ItemDataRole.UserRole)
        if annotation_id:
            self.selected_annotation_id = annotation_id
            self.update_canvas_marks()

    def rebuild_marker_list(self):
        self.marker_list.clear()
        for a in self.annotations:
            c = self.components[a.component_index]
            torque = self.torque_display(a)
            suffix = "  [僅標記]" if a.marker_only else ""
            text = f"{a.sequence}. {c.part_name}  {torque or NO_TORQUE}{suffix}"
            item = QListWidgetItem(text)
            item.setData(Qt.ItemDataRole.UserRole, a.annotation_id)
            self.marker_list.addItem(item)

    def update_canvas_marks(self):
        for item in self.annotation_items: self.canvas.scene.removeItem(item)
        self.annotation_items.clear()
        if not self.canvas.image_item: return
        rect = self.canvas.scene.sceneRect()
        for a in self.annotations:
            x, y, r = a.x*rect.width(), a.y*rect.height(), 18
            is_selected = a.annotation_id == self.selected_annotation_id
            pen_color = QColor("#0078d4") if is_selected else QColor("red")
            ellipse = self.canvas.scene.addEllipse(x-r, y-r, 2*r, 2*r, QPen(pen_color, 3 if not is_selected else 4))
            text = self.canvas.scene.addSimpleText(str(a.sequence)); text.setBrush(pen_color)
            font = text.font(); font.setBold(True); text.setFont(font)
            bounds = text.boundingRect(); text.setPos(x-bounds.width()/2, y-bounds.height()/2)
            self.annotation_items.extend([ellipse, text])

    def delete_last_annotation(self):
        if self.annotations:
            self.annotations.pop()
            self.rebuild_marker_list()
            self.update_canvas_marks()

    @staticmethod
    def torque_display(a):
        if a.torque is None or a.torque_unit == NO_TORQUE: return ""
        return LOCK_ONLY if a.torque_unit == LOCK_ONLY else f"{a.torque:g} {a.torque_unit}"

    def render_composite(self):
        base = Image.open(self.image_path).convert("RGBA")
        draw = ImageDraw.Draw(base); font = get_windows_font(max(16, min(base.size)//32))
        width, height = base.size
        for a in self.annotations:
            c = self.components[a.component_index]; x, y = int(a.x*width), int(a.y*height)
            r = max(12, min(width, height)//45)
            draw.ellipse((x-r, y-r, x+r, y+r), outline=MARKER_COLOR, width=max(2, r//7))
            sequence_text = str(a.sequence); bbox = draw.textbbox((0, 0), sequence_text, font=font)
            tw, th = bbox[2]-bbox[0], bbox[3]-bbox[1]
            draw.text((x-tw/2, y-th/2-bbox[1]), sequence_text, fill=MARKER_COLOR, font=font)
            if a.marker_only: continue
            label_x, label_y = x+r+8, y-r
            draw.text((label_x, label_y), c.part_name, fill=TEXT_COLOR, font=font)
            torque = self.torque_display(a)
            if torque:
                tb = draw.textbbox((0, 0), torque, font=font); pad = 5
                box = (label_x-pad, label_y+font.size+2, label_x+(tb[2]-tb[0])+pad, label_y+font.size+2+(tb[3]-tb[1])+pad*2)
                draw.rounded_rectangle(box, radius=4, outline=TORQUE_BOX_COLOR, width=1)
                draw.text((label_x, box[1]+pad), torque, fill=TEXT_COLOR, font=font)
        return base

    def update_preview(self):
        if not self.image_path: return
        try:
            image = self.render_composite(); path = Path(tempfile.gettempdir())/"sop_helper_preview.png"; image.save(path)
            self.preview_label.setPixmap(QPixmap(str(path)).scaled(self.preview_label.size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        except Exception as exc: self.status.setText(f"預覽失敗：{exc}")

    def save_png(self):
        path, _ = QFileDialog.getSaveFileName(self, "另存 PNG", "SOP_annotated.png", "PNG (*.png)")
        if path: self.render_composite().save(path); self.status.setText(f"已輸出：{path}")

    def save_pptx(self):
        if not self.image_path: QMessageBox.information(self, "提示", "請先載入主圖片"); return
        path, _ = QFileDialog.getSaveFileName(self, "另存 PPTX", "SOP_editable.pptx", "PowerPoint (*.pptx)")
        if path: export_pptx(path, self.image_path, self.components, self.annotations, self); self.status.setText(f"已輸出可編輯 PPTX：{path}")

    def save_project(self):
        path, _ = QFileDialog.getSaveFileName(self, "另存專案", "SOP_project.json", "JSON (*.json)")
        if path:
            data = {"project_version":"0.9.7", "source_image":self.image_path, "annotations":[asdict(a) for a in self.annotations], "components":[asdict(c) for c in self.components]}
            Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"); self.status.setText(f"已輸出專案：{path}")

    def copy_composite_image(self):
        if not win32clipboard: QMessageBox.critical(self, "缺少套件", "請安裝 pywin32"); return
        try:
            from io import BytesIO
            output = BytesIO(); self.render_composite().convert("RGB").save(output, "BMP"); data = output.getvalue()[14:]
            win32clipboard.OpenClipboard(); win32clipboard.EmptyClipboard(); win32clipboard.SetClipboardData(win32con.CF_DIB, data); win32clipboard.CloseClipboard()
            self.status.setText("合成圖片已複製到剪貼簿")
        except Exception as exc:
            try: win32clipboard.CloseClipboard()
            except Exception: pass
            QMessageBox.critical(self, "剪貼簿失敗", str(exc))

    def copy_pptx_shapes(self):
        if not self.image_path: QMessageBox.information(self, "提示", "請先載入主圖片"); return
        if win32com is None: QMessageBox.critical(self, "缺少套件", "請安裝 pywin32"); return
        temp_path = Path(tempfile.gettempdir()) / f"sop_helper_clip_{uuid.uuid4().hex}.pptx"
        powerpoint = presentation = None
        try:
            export_pptx(str(temp_path), self.image_path, self.components, self.annotations, self)
            powerpoint = win32com.client.Dispatch("PowerPoint.Application")
            presentation = powerpoint.Presentations.Open(str(temp_path), WithWindow=False)
            slide = presentation.Slides(1)
            shape_range = slide.Shapes.Range()
            if shape_range.Count == 0: raise RuntimeError("投影片中沒有可複製的物件")
            shape_range.Copy()
            self.status.setText("已複製 PowerPoint 物件到剪貼簿，請切換到目標簡報按 Ctrl+V")
        except Exception as exc:
            QMessageBox.critical(self, "複製 PowerPoint 物件失敗", f"{exc}\n\n請改用另存 PPTX 或複製合成圖片。")
        finally:
            try:
                if presentation is not None: presentation.Close()
            except Exception: pass
            try:
                if temp_path.exists(): temp_path.unlink()
            except Exception: pass


def get_windows_font(size):
    if not Path(WINDOWS_FONT_FILE).exists():
        raise RuntimeError(f"找不到微軟正黑體：{WINDOWS_FONT_FILE}")
    return ImageFont.truetype(WINDOWS_FONT_FILE, size=size)


def clean_value(value, default=None): return default if value is None or (isinstance(value, str) and not value.strip()) else value

def to_bool(value, default=False):
    if value is None or value == "": return default
    if isinstance(value, bool): return value
    return str(value).strip().lower() in {"true","1","yes","y","是","需要"}

def to_float(value): return None if value is None or value == "" else float(value)

def component_from_dict(row):
    return Component(str(clean_value(row.get("part_name"),"")).strip(), str(clean_value(row.get("part_number"),"")).strip(), str(clean_value(row.get("size"),"")).strip(), to_float(row.get("required_torque")), str(clean_value(row.get("torque_unit"),"")).strip(), str(clean_value(row.get("applicable_model"),"Universal")), str(clean_value(row.get("reference_image"),"")), to_bool(row.get("requires_oring")), str(clean_value(row.get("oring_spec"),"")), to_bool(row.get("requires_silicone_grease")), str(clean_value(row.get("silicone_grease_note"),"")), to_float(row.get("torque_min")), to_float(row.get("torque_max")), str(clean_value(row.get("notes"),"")))

def validate_components(components):
    errors, warnings, seen = [], [], set()
    for i, c in enumerate(components, 1):
        if not c.part_name: errors.append(f"第 {i} 列缺少 part_name")
        if not c.part_number: errors.append(f"第 {i} 列缺少 part_number")
        if not c.size: errors.append(f"第 {i} 列缺少 size")
        if c.part_number in seen: errors.append(f"重複 part_number：{c.part_number}")
        seen.add(c.part_number)
        target, unit = c.required_torque is not None, bool(c.torque_unit)
        if target != unit: errors.append(f"{c.part_number} 的 required_torque 與 torque_unit 必須同時填寫或同時留空")
        elif not target: warnings.append(f"{c.part_number}：未設定扭力，視為無需扭力元件")
        elif c.torque_min is not None and c.torque_max is not None and not (c.torque_min <= c.required_torque <= c.torque_max): errors.append(f"{c.part_number} 的扭力上下限不合理")
    if errors: raise ValueError("\n".join(errors))
    return warnings

def load_components(path):
    path = Path(path); rows=[]
    if path.suffix.lower()==".xlsx":
        wb=load_workbook(path,data_only=True,read_only=True); sheet=wb["components"] if "components" in wb.sheetnames else wb.active; headers=[c.value for c in next(sheet.iter_rows())]; rows=[dict(zip(headers,v)) for v in sheet.iter_rows(min_row=2,values_only=True) if not all(x is None for x in v)]
    elif path.suffix.lower() in {".yaml",".yml"}: rows=(yaml.safe_load(path.read_text(encoding="utf-8")) or {}).get("components",[])
    else: raise ValueError("只支援 .xlsx、.yaml、.yml")
    components=[component_from_dict(row) for row in rows]; return components, validate_components(components)

def add_textbox(slide, left, top, width, height, text, size=12, color=RGBColor(0, 0, 0), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text; run.font.name = WINDOWS_FONT_FAMILY; run.font.size = Pt(size); run.font.color.rgb = color
    return box

def fit_box_size(text, font_size=11, pad_x=Inches(.12), pad_y=Inches(.06), min_width=Inches(.5), min_height=Inches(.3)):
    char_width = int(Pt(font_size) * 0.62)
    text_width = char_width * max(len(text), 1)
    width = max(min_width, text_width + pad_x * 2)
    height = max(min_height, Pt(font_size) + pad_y * 2)
    return int(width), int(height)


def point_toward(from_x, from_y, to_x, to_y, back_off=0):
    dx, dy = to_x - from_x, to_y - from_y
    dist = (dx ** 2 + dy ** 2) ** 0.5
    if dist == 0:
        return int(to_x), int(to_y)
    ratio = (dist - back_off) / dist
    return int(from_x + dx * ratio), int(from_y + dy * ratio)


def add_arrowhead(connector_shape):
    ln = connector_shape.line._get_or_add_ln()
    for existing in ln.findall(qn('a:tailEnd')):
        ln.remove(existing)
    tail_end = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail_end)


def sequence_font_size(sequence, base_size=9):
    digit_count = len(str(sequence))
    if digit_count <= 1:
        return base_size
    elif digit_count == 2:
        return max(6, base_size - 2)
    else:
        return max(5, base_size - 3)


def export_pptx(path, image_path, components, annotations, window):
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); left,top=Inches(.5),Inches(.5)
    pic=slide.shapes.add_picture(image_path,left,top,width=Inches(8.5)); image_width=Inches(8.5); image_height=image_width*pic.height/pic.width
    for a in annotations:
        c=components[a.component_index]; x,y=left+int(image_width*a.x),top+int(image_height*a.y); radius=Inches(.10)
        circle=slide.shapes.add_shape(MSO_SHAPE.OVAL,x-radius,y-radius,radius*2,radius*2); circle.fill.background(); circle.line.color.rgb=RGBColor(220,0,0); circle.line.width=Pt(1.5)
        tf=circle.text_frame; tf.clear(); tf.word_wrap=False; tf.auto_size=None; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; run=p.add_run(); run.text=str(a.sequence); run.font.name=WINDOWS_FONT_FAMILY; run.font.size=Pt(sequence_font_size(a.sequence)); run.font.bold=True; run.font.color.rgb=RGBColor(220,0,0)
        tx,ty=x+Inches(.16),y-Inches(.18)
        if not a.marker_only:
            name_box=add_textbox(slide,tx,ty,Inches(2.8),Inches(.35),c.part_name,12,RGBColor(0,0,0))
            torque=window.torque_display(a)
            anchor_shape=name_box
            if torque:
                box_width,box_height=fit_box_size(torque,font_size=11)
                tb=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,tx,ty+Inches(.4),box_width,box_height); tb.fill.background(); tb.line.color.rgb=RGBColor(210,0,0); tb.line.width=Pt(0.75); tf=tb.text_frame; tf.clear(); tf.word_wrap=False; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.margin_top=0; tf.margin_bottom=0; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; run=p.add_run(); run.text=torque; run.font.name=WINDOWS_FONT_FAMILY; run.font.size=Pt(11); run.font.color.rgb=RGBColor(0,0,0)
                anchor_shape=tb
            anchor_left_mid_x, anchor_left_mid_y = anchor_shape.left, anchor_shape.top + anchor_shape.height // 2
            circle_right_mid_x, circle_right_mid_y = x + radius, y
            line=slide.shapes.add_connector(1, anchor_left_mid_x, anchor_left_mid_y, circle_right_mid_x, circle_right_mid_y)
            line.line.color.rgb=RGBColor(220,0,0); line.line.width=Pt(0.75)
            try:
                line.begin_connect(anchor_shape, 1)
                line.end_connect(circle, 3)
            except Exception:
                pass
            add_arrowhead(line)
    prs.save(path)

def main():
    app=QApplication(sys.argv); app.setFont(QFont(WINDOWS_FONT_FAMILY)); window=MainWindow(); window.show(); sys.exit(app.exec())

if __name__=="__main__": main()
