from __future__ import annotations

import json
import sys
import tempfile
from dataclasses import dataclass, asdict
from pathlib import Path

from PySide6.QtCore import QPointF, QRectF, Qt
from PySide6.QtGui import QColor, QPainter, QPen, QPixmap
from PySide6.QtWidgets import (
    QApplication,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QDoubleSpinBox,
    QSplitter,
    QVBoxLayout,
    QCheckBox,
)
from PySide6.QtWidgets import QGraphicsView, QGraphicsScene, QWidget

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import yaml
except ImportError:
    yaml = None

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ImportError:
    Presentation = None

try:
    from PIL import Image, ImageDraw
except ImportError:
    Image = None

try:
    import win32clipboard
    import win32con
except ImportError:
    win32clipboard = None


APP_TITLE = "SOP_Helper MVP v0.2"
NO_TORQUE_TEXT = "無需扭力"


@dataclass
class Component:
    part_name: str
    part_number: str
    size: str
    required_torque: float | None
    torque_unit: str
    applicable_model: str = "Universal"
    reference_image: str = ""
    requires_oring: bool = False
    oring_spec: str = ""
    requires_silicone_grease: bool = False
    silicone_grease_note: str = ""
    torque_min: float | None = None
    torque_max: float | None = None
    notes: str = ""

    @property
    def has_torque_spec(self) -> bool:
        return self.required_torque is not None and bool(self.torque_unit)


@dataclass
class Annotation:
    annotation_id: str
    component_index: int
    x: float
    y: float
    sequence: int
    torque: float | None
    torque_unit: str
    torque_overridden: bool
    override_reason: str
    label_x: float
    label_y: float


class AnnotationDialog(QDialog):
    def __init__(self, component: Component, sequence: int, parent=None):
        super().__init__(parent)
        self.setWindowTitle("設定標記")
        self.component = component

        form = QFormLayout(self)
        self.component_label = QLabel(
            f"{component.part_name} ({component.part_number})\n尺寸：{component.size}"
        )
        form.addRow("料件", self.component_label)

        self.has_torque = component.has_torque_spec
        self.use_custom = QCheckBox("使用自訂扭力（與規格不同）")
        self.use_custom.setChecked(False)
        self.use_custom.toggled.connect(self._toggle_custom)
        form.addRow(self.use_custom)

        self.torque = QDoubleSpinBox()
        self.torque.setRange(0, 1_000_000)
        self.torque.setDecimals(4)
        self.unit = QLineEdit()

        if self.has_torque:
            self.torque.setValue(component.required_torque)
            self.unit.setText(component.torque_unit)
            self.torque.setEnabled(False)
            self.unit.setEnabled(False)
        else:
            self.torque.setValue(0)
            self.unit.setText("")
            self.torque.setEnabled(False)
            self.unit.setEnabled(False)
            self.use_custom.setText("此料件未設定扭力規格，勾選以手動輸入本次扭力（選填）")

        form.addRow("本次扭力", self.torque)
        form.addRow("單位", self.unit)

        self.reason = QLineEdit()
        self.reason.setPlaceholderText("若使用自訂扭力，請填寫原因")
        self.reason.setEnabled(False)
        form.addRow("自訂原因", self.reason)

        self.sequence = QSpinBox()
        self.sequence.setRange(1, 999)
        self.sequence.setValue(sequence)
        form.addRow("鎖附順序", self.sequence)

        torque_hint = (
            f"{component.required_torque:g} {component.torque_unit}"
            if self.has_torque
            else NO_TORQUE_TEXT
        )
        self.note = QLabel(
            f"規格扭力：{torque_hint}\n"
            f"機型：{component.applicable_model}\n"
            f"O-ring：{component.oring_spec if component.requires_oring else '無'}\n"
            f"圖矽油：{component.silicone_grease_note if component.requires_silicone_grease else '無'}"
        )
        form.addRow("規格提示", self.note)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok
            | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(self._on_accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def _toggle_custom(self, checked: bool):
        self.torque.setEnabled(checked)
        self.unit.setEnabled(checked)
        self.reason.setEnabled(checked)
        if not checked and self.has_torque:
            self.torque.setValue(self.component.required_torque)
            self.unit.setText(self.component.torque_unit)

    def _on_accept(self):
        if self.use_custom.isChecked() and not self.reason.text().strip():
            QMessageBox.warning(self, "缺少原因", "使用自訂扭力時請填寫原因")
            return
        self.accept()

    def values(self):
        overridden = self.use_custom.isChecked()
        if overridden:
            torque = self.torque.value()
            unit = self.unit.text().strip()
        elif self.has_torque:
            torque = self.component.required_torque
            unit = self.component.torque_unit
        else:
            torque = None
            unit = ""
        reason = self.reason.text().strip() if overridden else ""
        return torque, unit, self.sequence.value(), overridden, reason


class AnnotationCanvas(QGraphicsView):
    def __init__(self, owner, parent=None):
        super().__init__(parent)
        self.owner = owner
        self.scene = QGraphicsScene(self)
        self.setScene(self.scene)
        self.setRenderHints(QPainter.RenderHint.Antialiasing)
        self.setDragMode(QGraphicsView.DragMode.ScrollHandDrag)
        self.setBackgroundBrush(QColor("#eeeeee"))
        self.image_item = None
        self.pixmap = None

    def set_image(self, path: str):
        self.scene.clear()
        self.owner.annotation_items.clear()
        self.image_item = None
        self.pixmap = QPixmap(path)
        if self.pixmap.isNull():
            raise ValueError("無法讀取圖片")
        self.image_item = self.scene.addPixmap(self.pixmap)
        self.image_item.setZValue(0)
        self.scene.setSceneRect(QRectF(self.pixmap.rect()))
        self.fitInView(self.scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.image_item:
            self.fitInView(self.scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self.image_item:
            scene_pos = self.mapToScene(event.position().toPoint())
            if self.scene.sceneRect().contains(scene_pos):
                self.owner.add_annotation_at(scene_pos)
                return
        super().mousePressEvent(event)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(1500, 850)
        self.components: list[Component] = []
        self.annotations: list[Annotation] = []
        self.annotation_items = []
        self.image_path = ""

        self.canvas = AnnotationCanvas(self)
        self.component_list = QListWidget()
        self.component_list.itemDoubleClicked.connect(self.add_selected_component)

        self.status = QLabel("請先載入圖片與 Excel/YAML 規格")
        self.preview_label = QLabel("輸出預覽尚未產生")
        self.preview_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_label.setMinimumWidth(320)
        self.preview_label.setStyleSheet("background:#eeeeee; color:#555;")

        left = self.build_left_panel()
        right = self.build_right_panel()
        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(left)
        splitter.addWidget(self.canvas)
        splitter.addWidget(right)
        splitter.setSizes([300, 800, 380])
        self.setCentralWidget(splitter)
        self.statusBar().addWidget(self.status)

    def build_left_panel(self):
        box = QGroupBox("資料與料件")
        layout = QVBoxLayout(box)

        load_image = QPushButton("載入主圖片")
        load_image.clicked.connect(self.load_image)
        layout.addWidget(load_image)

        load_spec = QPushButton("載入 Excel / YAML")
        load_spec.clicked.connect(self.load_spec)
        layout.addWidget(load_spec)

        layout.addWidget(QLabel("可用料件（雙擊後在圖片上新增標記）"))
        layout.addWidget(self.component_list, 1)

        add_button = QPushButton("使用選取料件新增標記")
        add_button.clicked.connect(self.add_selected_component)
        layout.addWidget(add_button)

        delete_button = QPushButton("刪除最後一個標記")
        delete_button.clicked.connect(self.delete_last_annotation)
        layout.addWidget(delete_button)
        return box

    def build_right_panel(self):
        box = QGroupBox("輸出與預覽")
        layout = QVBoxLayout(box)
        layout.addWidget(self.preview_label, 1)

        refresh = QPushButton("更新輸出預覽")
        refresh.clicked.connect(self.update_preview)
        layout.addWidget(refresh)

        copy_image = QPushButton("複製合成圖片")
        copy_image.clicked.connect(self.copy_composite_image)
        layout.addWidget(copy_image)

        save_image = QPushButton("另存 PNG")
        save_image.clicked.connect(self.save_png)
        layout.addWidget(save_image)

        save_pptx = QPushButton("另存可編輯 PPTX")
        save_pptx.clicked.connect(self.save_pptx)
        layout.addWidget(save_pptx)

        save_project = QPushButton("另存專案 JSON")
        save_project.clicked.connect(self.save_project)
        layout.addWidget(save_project)
        return box

    def load_image(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "選擇主圖片", "", "Images (*.png *.jpg *.jpeg *.bmp)"
        )
        if not path:
            return
        try:
            self.canvas.set_image(path)
            self.image_path = path
            self.status.setText(f"已載入圖片：{Path(path).name}")
            self.update_preview()
        except Exception as exc:
            QMessageBox.critical(self, "載入失敗", str(exc))

    def load_spec(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "選擇規格檔", "", "Excel/YAML (*.xlsx *.yaml *.yml)"
        )
        if not path:
            return
        try:
            self.components, warnings = load_components(path)
            self.component_list.clear()
            for index, component in enumerate(self.components):
                torque = (
                    f"{component.required_torque:g} {component.torque_unit}"
                    if component.has_torque_spec
                    else NO_TORQUE_TEXT
                )
                item = QListWidgetItem(
                    f"{component.part_name} | {component.part_number} | {torque}"
                )
                item.setData(Qt.ItemDataRole.UserRole, index)
                self.component_list.addItem(item)
            msg = f"已載入 {len(self.components)} 筆料件規格"
            if warnings:
                msg += f"（{len(warnings)} 筆提示，詳見下方視窗）"
                QMessageBox.information(self, "載入完成，附帶提示", "\n".join(warnings))
            self.status.setText(msg)
        except Exception as exc:
            QMessageBox.critical(self, "規格載入失敗", str(exc))

    def selected_component_index(self):
        item = self.component_list.currentItem()
        if not item:
            return None
        return item.data(Qt.ItemDataRole.UserRole)

    def add_selected_component(self):
        index = self.selected_component_index()
        if index is None:
            QMessageBox.information(self, "提示", "請先選擇料件")
            return
        if not self.canvas.image_item:
            QMessageBox.information(self, "提示", "請先載入主圖片")
            return
        center = self.canvas.scene.sceneRect().center()
        self.add_annotation_at(center, index)

    def add_annotation_at(self, scene_pos: QPointF, component_index=None):
        if component_index is None:
            component_index = self.selected_component_index()
        if component_index is None:
            QMessageBox.information(self, "提示", "請先在左側選擇料件")
            return
        component = self.components[component_index]
        dialog = AnnotationDialog(component, len(self.annotations) + 1, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        torque, unit, sequence, overridden, reason = dialog.values()
        rect = self.canvas.scene.sceneRect()
        annotation = Annotation(
            annotation_id=f"F{len(self.annotations) + 1:03d}",
            component_index=component_index,
            x=scene_pos.x() / rect.width(),
            y=scene_pos.y() / rect.height(),
            sequence=sequence,
            torque=torque,
            torque_unit=unit,
            torque_overridden=overridden,
            override_reason=reason,
            label_x=min(max(scene_pos.x() / rect.width() - 0.1, 0), 0.8),
            label_y=min(max(scene_pos.y() / rect.height() - 0.15, 0), 0.8),
        )
        self.annotations.append(annotation)
        self.update_canvas_marks()
        self.update_preview()

    def update_canvas_marks(self):
        for item in self.annotation_items:
            self.canvas.scene.removeItem(item)
        self.annotation_items.clear()
        if not self.canvas.image_item:
            return
        rect = self.canvas.scene.sceneRect()
        for annotation in self.annotations:
            x = annotation.x * rect.width()
            y = annotation.y * rect.height()
            r = 14
            color = QColor("orange") if annotation.torque_overridden else QColor("red")
            ellipse = self.canvas.scene.addEllipse(
                x - r, y - r, 2 * r, 2 * r, QPen(color, 3)
            )
            ellipse.setZValue(2)
            text = self.canvas.scene.addSimpleText(str(annotation.sequence))
            text.setBrush(color)
            text.setPos(x - 5, y - 10)
            text.setZValue(3)
            self.annotation_items.extend([ellipse, text])

    def delete_last_annotation(self):
        if self.annotations:
            self.annotations.pop()
            self.update_canvas_marks()
            self.update_preview()

    def torque_display(self, annotation: Annotation) -> str:
        if annotation.torque is None:
            return NO_TORQUE_TEXT
        base = f"{annotation.torque:g} {annotation.torque_unit}"
        if annotation.torque_overridden:
            base += "（自訂）"
        return base

    def render_composite(self):
        if not self.image_path:
            raise ValueError("尚未載入主圖片")
        base = Image.open(self.image_path).convert("RGBA")
        draw = ImageDraw.Draw(base)
        width, height = base.size
        for annotation in self.annotations:
            component = self.components[annotation.component_index]
            x = int(annotation.x * width)
            y = int(annotation.y * height)
            r = max(12, min(width, height) // 45)
            color = (230, 140, 0, 255) if annotation.torque_overridden else (220, 0, 0, 255)
            draw.ellipse((x - r, y - r, x + r, y + r), outline=color, width=max(3, r // 5))
            draw.text((x + r + 5, y - r), f"{annotation.sequence}. {component.part_name}", fill=color)
            draw.text((x + r + 5, y + 3), self.torque_display(annotation), fill=color)
        return base

    def update_preview(self):
        if not self.image_path or Image is None:
            return
        try:
            image = self.render_composite()
            path = Path(tempfile.gettempdir()) / "sop_helper_preview.png"
            image.save(path)
            pixmap = QPixmap(str(path))
            self.preview_label.setPixmap(
                pixmap.scaled(
                    self.preview_label.size(),
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
            )
        except Exception as exc:
            self.status.setText(f"預覽失敗：{exc}")

    def save_png(self):
        if Image is None:
            QMessageBox.critical(self, "缺少套件", "請安裝 Pillow")
            return
        path, _ = QFileDialog.getSaveFileName(self, "另存 PNG", "SOP_annotated.png", "PNG (*.png)")
        if not path:
            return
        try:
            self.render_composite().save(path)
            self.status.setText(f"已輸出：{path}")
        except Exception as exc:
            QMessageBox.critical(self, "輸出失敗", str(exc))

    def save_pptx(self):
        if Presentation is None:
            QMessageBox.critical(self, "缺少套件", "請安裝 python-pptx")
            return
        if not self.image_path:
            QMessageBox.information(self, "提示", "請先載入主圖片")
            return
        path, _ = QFileDialog.getSaveFileName(self, "另存 PPTX", "SOP_editable.pptx", "PowerPoint (*.pptx)")
        if not path:
            return
        try:
            export_pptx(path, self.image_path, self.components, self.annotations, self)
            self.status.setText(f"已輸出可編輯 PPTX：{path}")
        except Exception as exc:
            QMessageBox.critical(self, "PPTX 輸出失敗", str(exc))

    def save_project(self):
        path, _ = QFileDialog.getSaveFileName(self, "另存專案", "SOP_project.json", "JSON (*.json)")
        if not path:
            return
        data = {
            "project_version": "0.2",
            "source_image": self.image_path,
            "annotations": [asdict(a) for a in self.annotations],
            "components": [asdict(c) for c in self.components],
        }
        Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        self.status.setText(f"已輸出專案：{path}")

    def copy_composite_image(self):
        if Image is None:
            QMessageBox.critical(self, "缺少套件", "請安裝 Pillow")
            return
        if win32clipboard is None:
            QMessageBox.critical(self, "缺少套件", "請安裝 pywin32")
            return
        try:
            image = self.render_composite().convert("RGB")
            from io import BytesIO
            output = BytesIO()
            image.save(output, "BMP")
            data = output.getvalue()[14:]
            win32clipboard.OpenClipboard()
            win32clipboard.EmptyClipboard()
            win32clipboard.SetClipboardData(win32con.CF_DIB, data)
            win32clipboard.CloseClipboard()
            self.status.setText("合成圖片已複製到剪貼簿")
        except Exception as exc:
            try:
                win32clipboard.CloseClipboard()
            except Exception:
                pass
            QMessageBox.critical(self, "剪貼簿失敗", str(exc))


def clean_value(value, default=None):
    if value is None:
        return default
    if isinstance(value, str) and not value.strip():
        return default
    return value


def to_bool(value, default=False):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "1", "yes", "y", "是", "需要"}


def to_float(value):
    if value is None or value == "":
        return None
    return float(value)


def component_from_dict(row):
    torque = row.get("required_torque")
    return Component(
        part_name=str(clean_value(row.get("part_name"), "")).strip(),
        part_number=str(clean_value(row.get("part_number"), "")).strip(),
        size=str(clean_value(row.get("size"), "")).strip(),
        required_torque=to_float(torque),
        torque_unit=str(clean_value(row.get("torque_unit"), "")).strip(),
        applicable_model=str(clean_value(row.get("applicable_model"), "Universal")),
        reference_image=str(clean_value(row.get("reference_image"), "")),
        requires_oring=to_bool(row.get("requires_oring"), False),
        oring_spec=str(clean_value(row.get("oring_spec"), "")),
        requires_silicone_grease=to_bool(row.get("requires_silicone_grease"), False),
        silicone_grease_note=str(clean_value(row.get("silicone_grease_note"), "")),
        torque_min=to_float(row.get("torque_min")),
        torque_max=to_float(row.get("torque_max")),
        notes=str(clean_value(row.get("notes"), "")),
    )


def validate_components(components):
    errors = []
    warnings = []
    seen = set()
    for i, c in enumerate(components, 1):
        if not c.part_name:
            errors.append(f"第 {i} 列缺少 part_name")
        if not c.part_number:
            errors.append(f"第 {i} 列缺少 part_number")
        if not c.size:
            errors.append(f"第 {i} 列缺少 size")
        if c.part_number in seen:
            errors.append(f"重複 part_number：{c.part_number}")
        seen.add(c.part_number)

        has_target = c.required_torque is not None
        has_unit = bool(c.torque_unit)
        if has_target != has_unit:
            errors.append(
                f"{c.part_number} 的 required_torque 與 torque_unit 必須同時填寫或同時留空"
            )
        elif not has_target and not has_unit:
            warnings.append(f"{c.part_number}：未設定扭力，視為非鎖螺絲元件（例如 O-ring／墊片）")
        else:
            if c.torque_min is not None and c.torque_max is not None:
                if not (c.torque_min <= c.required_torque <= c.torque_max):
                    errors.append(f"{c.part_number} 的扭力上下限不合理")
    if errors:
        raise ValueError("\n".join(errors))
    return warnings


def load_components(path):
    path = Path(path)
    rows = []
    if path.suffix.lower() == ".xlsx":
        if load_workbook is None:
            raise RuntimeError("請安裝 openpyxl")
        wb = load_workbook(path, data_only=True, read_only=True)
        sheet = wb["components"] if "components" in wb.sheetnames else wb.active
        headers = [cell.value for cell in next(sheet.iter_rows())]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in values):
                continue
            rows.append(dict(zip(headers, values)))
    elif path.suffix.lower() in {".yaml", ".yml"}:
        if yaml is None:
            raise RuntimeError("請安裝 PyYAML")
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        rows = data.get("components", [])
    else:
        raise ValueError("只支援 .xlsx、.yaml、.yml")
    components = [component_from_dict(row) for row in rows]
    warnings = validate_components(components)
    return components, warnings


def export_pptx(path, image_path, components, annotations, window: MainWindow):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    left = Inches(0.5)
    top = Inches(0.5)
    pic = slide.shapes.add_picture(image_path, left, top, width=Inches(8.5))
    image_width = Inches(8.5)
    image_height = image_width * pic.height / pic.width
    for annotation in annotations:
        component = components[annotation.component_index]
        x = left + int(image_width * annotation.x)
        y = top + int(image_height * annotation.y)
        radius = Inches(0.12)
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, x - radius, y - radius, radius * 2, radius * 2
        )
        circle.fill.background()
        circle_color = RGBColor(230, 140, 0) if annotation.torque_overridden else RGBColor(220, 0, 0)
        circle.line.color.rgb = circle_color
        circle.line.width = Pt(2.5)

        tx = x + Inches(0.16)
        ty = y - Inches(0.18)
        textbox = slide.shapes.add_textbox(tx, ty, Inches(2.5), Inches(0.5))
        tf = textbox.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        torque_text = window.torque_display(annotation)
        run.text = f"{annotation.sequence}. {component.part_name}\n{torque_text}"
        run.font.size = Pt(12)
        run.font.color.rgb = circle_color

        line = slide.shapes.add_connector(1, tx, ty + Inches(0.22), x, y)
        line.line.color.rgb = circle_color
        line.line.width = Pt(2)
    prs.save(path)


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
