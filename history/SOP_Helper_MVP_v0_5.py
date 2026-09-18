from __future__ import annotations

import json
import sys
import tempfile
from dataclasses import dataclass, asdict
from pathlib import Path

from PySide6.QtCore import QPointF, QRectF, Qt
from PySide6.QtGui import QColor, QPainter, QPen, QPixmap
from PySide6.QtWidgets import (
    QApplication, QComboBox, QDialog, QDialogButtonBox, QFileDialog,
    QFormLayout, QGroupBox, QLabel, QListWidget, QListWidgetItem,
    QMainWindow, QMessageBox, QPushButton, QSpinBox, QDoubleSpinBox,
    QSplitter, QVBoxLayout, QGraphicsView, QGraphicsScene
)
from openpyxl import load_workbook
import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

try:
    import win32clipboard
    import win32con
except ImportError:
    win32clipboard = None
    win32con = None

APP_TITLE = "SOP_Helper MVP v0.5"
TORQUE_UNITS = ["kgf-cm", "N-M", "無需扭力", "鎖緊就好"]
NO_TORQUE = "無需扭力"
LOCK_ONLY = "鎖緊就好"
TEXT_COLOR = (0, 0, 0, 255)
TORQUE_BOX_COLOR = (210, 0, 0, 255)
MARKER_COLOR = (220, 0, 0, 255)


@dataclass
class Component:
    part_name: str
    part_number: str
    size: str
    required_torque: float | None
    torque_unit: str
    applicable_model: str = "Universal"
    reference_image: str = ""
    requires_oring: bool = False
    oring_spec: str = ""
    requires_silicone_grease: bool = False
    silicone_grease_note: str = ""
    torque_min: float | None = None
    torque_max: float | None = None
    notes: str = ""

    @property
    def has_torque_spec(self):
        return self.required_torque is not None and bool(self.torque_unit)


@dataclass
class Annotation:
    annotation_id: str
    component_index: int
    x: float
    y: float
    sequence: int
    torque: float | None
    torque_unit: str
    label_x: float
    label_y: float


class AnnotationDialog(QDialog):
    def __init__(self, component, default_sequence, used_sequences, parent=None):
        super().__init__(parent)
        self.setWindowTitle("設定標記")
        self.component = component
        self.used_sequences = set(used_sequences)
        form = QFormLayout(self)
        form.addRow("料件", QLabel(f"{component.part_name} ({component.part_number})\n尺寸：{component.size}"))
        self.torque = QDoubleSpinBox()
        self.torque.setRange(0, 1_000_000)
        self.torque.setDecimals(4)
        if component.required_torque is not None:
            self.torque.setValue(component.required_torque)
        form.addRow("本次扭力", self.torque)
        self.unit = QComboBox()
        self.unit.addItems(TORQUE_UNITS)
        if component.has_torque_spec and component.torque_unit in TORQUE_UNITS:
            self.unit.setCurrentText(component.torque_unit)
        elif not component.has_torque_spec:
            self.unit.setCurrentText(NO_TORQUE)
            self.torque.setEnabled(False)
        form.addRow("單位", self.unit)
        self.unit.currentTextChanged.connect(self.on_unit_changed)
        self.torque.valueChanged.connect(self.on_torque_changed)
        self.sequence = QSpinBox()
        self.sequence.setRange(1, 999)
        self.sequence.setValue(default_sequence)
        form.addRow("標記順序", self.sequence)
        hint = f"{component.required_torque:g} {component.torque_unit}" if component.has_torque_spec else NO_TORQUE
        form.addRow("規格提示", QLabel(
            f"規格扭力：{hint}\n機型：{component.applicable_model}\n"
            f"O-ring：{component.oring_spec if component.requires_oring else '無'}\n"
            f"塗矽油：{component.silicone_grease_note if component.requires_silicone_grease else '無'}"
        ))
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.validate_and_accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def on_unit_changed(self, unit):
        numeric = unit in {"kgf-cm", "N-M"}
        self.torque.setEnabled(numeric)
        if unit in {NO_TORQUE, LOCK_ONLY}:
            self.torque.setValue(0)

    def on_torque_changed(self, value):
        if value == 0 and self.unit.currentText() in {"kgf-cm", "N-M"}:
            self.unit.blockSignals(True)
            self.unit.setCurrentText(NO_TORQUE)
            self.unit.blockSignals(False)
            self.torque.setEnabled(False)

    def validate_and_accept(self):
        sequence = self.sequence.value()
        if sequence in self.used_sequences:
            QMessageBox.warning(self, "標記順序重複", f"標記順序 {sequence} 已存在，請改用其他數字。")
            return
        self.accept()

    def values(self):
        unit = self.unit.currentText()
        if unit == NO_TORQUE or self.torque.value() == 0:
            return None, NO_TORQUE, self.sequence.value()
        return self.torque.value(), unit, self.sequence.value()


class AnnotationCanvas(QGraphicsView):
    def __init__(self, owner):
        super().__init__()
        self.owner = owner
        self.scene = QGraphicsScene(self)
        self.setScene(self.scene)
        self.setRenderHints(QPainter.RenderHint.Antialiasing)
        self.setDragMode(QGraphicsView.DragMode.ScrollHandDrag)
        self.setBackgroundBrush(QColor("#eeeeee"))
        self.image_item = None

    def set_image(self, path):
        self.scene.clear()
        self.owner.annotation_items.clear()
        self.image_item = None
        pixmap = QPixmap(path)
        if pixmap.isNull():
            raise ValueError("無法讀取圖片")
        self.image_item = self.scene.addPixmap(pixmap)
        self.scene.setSceneRect(QRectF(pixmap.rect()))
        self.fitInView(self.scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.image_item:
            self.fitInView(self.scene.sceneRect(), Qt.AspectRatioMode.KeepAspectRatio)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self.image_item:
            pos = self.mapToScene(event.position().toPoint())
            if self.scene.sceneRect().contains(pos):
                self.owner.add_annotation_at(pos)
                return
        super().mousePressEvent(event)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(1500, 850)
        self.components = []
        self.annotations = []
        self.annotation_items = []
        self.image_path = ""
        self.canvas = AnnotationCanvas(self)
        self.component_list = QListWidget()
        self.component_list.itemDoubleClicked.connect(self.add_selected_component)
        self.status = QLabel("請先載入圖片與 Excel/YAML 規格")
        self.preview_label = QLabel("輸出預覽尚未產生")
        self.preview_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_label.setMinimumWidth(320)
        self.preview_label.setStyleSheet("background:#eeeeee; color:#555;")
        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(self.build_left_panel())
        splitter.addWidget(self.canvas)
        splitter.addWidget(self.build_right_panel())
        splitter.setSizes([300, 800, 380])
        self.setCentralWidget(splitter)
        self.statusBar().addWidget(self.status)

    def build_left_panel(self):
        box = QGroupBox("資料與料件")
        layout = QVBoxLayout(box)
        for text, slot in [("載入主圖片", self.load_image), ("載入 Excel / YAML", self.load_spec)]:
            b = QPushButton(text); b.clicked.connect(slot); layout.addWidget(b)
        layout.addWidget(QLabel("可用料件（雙擊後在圖片上新增標記）"))
        layout.addWidget(self.component_list, 1)
        b = QPushButton("使用選取料件新增標記"); b.clicked.connect(self.add_selected_component); layout.addWidget(b)
        b = QPushButton("刪除最後一個標記"); b.clicked.connect(self.delete_last_annotation); layout.addWidget(b)
        return box

    def build_right_panel(self):
        box = QGroupBox("輸出與預覽")
        layout = QVBoxLayout(box)
        layout.addWidget(self.preview_label, 1)
        for text, slot in [("更新輸出預覽", self.update_preview), ("複製合成圖片", self.copy_composite_image), ("另存 PNG", self.save_png), ("另存可編輯 PPTX", self.save_pptx), ("另存專案 JSON", self.save_project)]:
            b = QPushButton(text); b.clicked.connect(slot); layout.addWidget(b)
        return box

    def load_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "選擇主圖片", "", "Images (*.png *.jpg *.jpeg *.bmp)")
        if not path: return
        try:
            self.canvas.set_image(path); self.image_path = path
            self.status.setText(f"已載入圖片：{Path(path).name}"); self.update_preview()
        except Exception as exc: QMessageBox.critical(self, "載入失敗", str(exc))

    def load_spec(self):
        path, _ = QFileDialog.getOpenFileName(self, "選擇規格檔", "", "Excel/YAML (*.xlsx *.yaml *.yml)")
        if not path: return
        try:
            self.components, warnings = load_components(path)
            self.component_list.clear()
            for index, c in enumerate(self.components):
                torque = f"{c.required_torque:g} {c.torque_unit}" if c.has_torque_spec else NO_TORQUE
                item = QListWidgetItem(f"{c.part_name} | {c.part_number} | {torque}")
                item.setData(Qt.ItemDataRole.UserRole, index); self.component_list.addItem(item)
            if warnings: QMessageBox.information(self, "載入完成，附帶提示", "\n".join(warnings))
            self.status.setText(f"已載入 {len(self.components)} 筆料件規格")
        except Exception as exc: QMessageBox.critical(self, "規格載入失敗", str(exc))

    def selected_component_index(self):
        item = self.component_list.currentItem()
        return item.data(Qt.ItemDataRole.UserRole) if item else None

    def next_sequence(self):
        return max((a.sequence for a in self.annotations), default=0) + 1

    def add_selected_component(self):
        index = self.selected_component_index()
        if index is None: QMessageBox.information(self, "提示", "請先選擇料件"); return
        if not self.canvas.image_item: QMessageBox.information(self, "提示", "請先載入主圖片"); return
        self.add_annotation_at(self.canvas.scene.sceneRect().center(), index)

    def add_annotation_at(self, pos, component_index=None):
        if component_index is None: component_index = self.selected_component_index()
        if component_index is None: QMessageBox.information(self, "提示", "請先在左側選擇料件"); return
        dialog = AnnotationDialog(self.components[component_index], self.next_sequence(), [a.sequence for a in self.annotations], self)
        if dialog.exec() != QDialog.DialogCode.Accepted: return
        torque, unit, sequence = dialog.values()
        rect = self.canvas.scene.sceneRect()
        self.annotations.append(Annotation(f"F{len(self.annotations)+1:03d}", component_index, pos.x()/rect.width(), pos.y()/rect.height(), sequence, torque, unit, 0, 0))
        self.update_canvas_marks(); self.update_preview()

    def update_canvas_marks(self):
        for item in self.annotation_items: self.canvas.scene.removeItem(item)
        self.annotation_items.clear()
        if not self.canvas.image_item: return
        rect = self.canvas.scene.sceneRect()
        for a in self.annotations:
            x, y, r = a.x*rect.width(), a.y*rect.height(), 18
            ellipse = self.canvas.scene.addEllipse(x-r, y-r, 2*r, 2*r, QPen(QColor("red"), 3))
            text = self.canvas.scene.addSimpleText(str(a.sequence)); text.setBrush(QColor("red"))
            font = text.font(); font.setBold(True); text.setFont(font)
            bounds = text.boundingRect(); text.setPos(x-bounds.width()/2, y-bounds.height()/2)
            self.annotation_items.extend([ellipse, text])

    def delete_last_annotation(self):
        if self.annotations: self.annotations.pop(); self.update_canvas_marks(); self.update_preview()

    @staticmethod
    def torque_display(a):
        if a.torque is None or a.torque_unit == NO_TORQUE: return ""
        return LOCK_ONLY if a.torque_unit == LOCK_ONLY else f"{a.torque:g} {a.torque_unit}"

    def render_composite(self):
        base = Image.open(self.image_path).convert("RGBA")
        draw = ImageDraw.Draw(base); font = get_windows_font(max(16, min(base.size)//32))
        width, height = base.size
        for a in self.annotations:
            c = self.components[a.component_index]; x, y = int(a.x*width), int(a.y*height)
            r = max(12, min(width, height)//45)
            draw.ellipse((x-r, y-r, x+r, y+r), outline=MARKER_COLOR, width=max(3, r//5))
            # 順序只放在圓圈中心，不再放在料件名稱前。
            sequence_text = str(a.sequence)
            bbox = draw.textbbox((0, 0), sequence_text, font=font)
            tw, th = bbox[2]-bbox[0], bbox[3]-bbox[1]
            draw.text((x-tw/2, y-th/2-bbox[1]), sequence_text, fill=MARKER_COLOR, font=font)
            label_x, label_y = x+r+8, y-r
            draw.text((label_x, label_y), c.part_name, fill=TEXT_COLOR, font=font)
            torque = self.torque_display(a)
            if torque:
                tb = draw.textbbox((0, 0), torque, font=font)
                pad = 6
                box = (label_x-pad, label_y+font.size+2, label_x+(tb[2]-tb[0])+pad, label_y+font.size+2+(tb[3]-tb[1])+pad*2)
                draw.rounded_rectangle(box, radius=5, outline=TORQUE_BOX_COLOR, width=3)
                draw.text((label_x, box[1]+pad), torque, fill=TEXT_COLOR, font=font)
        return base

    def update_preview(self):
        if not self.image_path: return
        try:
            image = self.render_composite(); path = Path(tempfile.gettempdir())/"sop_helper_preview.png"; image.save(path)
            self.preview_label.setPixmap(QPixmap(str(path)).scaled(self.preview_label.size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        except Exception as exc: self.status.setText(f"預覽失敗：{exc}")

    def save_png(self):
        path, _ = QFileDialog.getSaveFileName(self, "另存 PNG", "SOP_annotated.png", "PNG (*.png)")
        if path: self.render_composite().save(path); self.status.setText(f"已輸出：{path}")

    def save_pptx(self):
        if not self.image_path: QMessageBox.information(self, "提示", "請先載入主圖片"); return
        path, _ = QFileDialog.getSaveFileName(self, "另存 PPTX", "SOP_editable.pptx", "PowerPoint (*.pptx)")
        if path: export_pptx(path, self.image_path, self.components, self.annotations, self); self.status.setText(f"已輸出可編輯 PPTX：{path}")

    def save_project(self):
        path, _ = QFileDialog.getSaveFileName(self, "另存專案", "SOP_project.json", "JSON (*.json)")
        if path:
            data = {"project_version":"0.5", "source_image":self.image_path, "annotations":[asdict(a) for a in self.annotations], "components":[asdict(c) for c in self.components]}
            Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"); self.status.setText(f"已輸出專案：{path}")

    def copy_composite_image(self):
        if not win32clipboard: QMessageBox.critical(self, "缺少套件", "請安裝 pywin32"); return
        try:
            from io import BytesIO
            output = BytesIO(); self.render_composite().convert("RGB").save(output, "BMP"); data = output.getvalue()[14:]
            win32clipboard.OpenClipboard(); win32clipboard.EmptyClipboard(); win32clipboard.SetClipboardData(win32con.CF_DIB, data); win32clipboard.CloseClipboard()
            self.status.setText("合成圖片已複製到剪貼簿")
        except Exception as exc:
            try: win32clipboard.CloseClipboard()
            except Exception: pass
            QMessageBox.critical(self, "剪貼簿失敗", str(exc))


def get_windows_font(size):
    for path in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc", r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simsun.ttc", r"C:\Windows\Fonts\arial.ttf"]:
        if Path(path).exists():
            try: return ImageFont.truetype(path, size=size)
            except OSError: pass
    raise RuntimeError("找不到可顯示中文的 Windows 字型")


def clean_value(value, default=None): return default if value is None or (isinstance(value, str) and not value.strip()) else value

def to_bool(value, default=False):
    if value is None or value == "": return default
    if isinstance(value, bool): return value
    return str(value).strip().lower() in {"true","1","yes","y","是","需要"}

def to_float(value): return None if value is None or value == "" else float(value)

def component_from_dict(row):
    return Component(str(clean_value(row.get("part_name"),"")).strip(), str(clean_value(row.get("part_number"),"")).strip(), str(clean_value(row.get("size"),"")).strip(), to_float(row.get("required_torque")), str(clean_value(row.get("torque_unit"),"")).strip(), str(clean_value(row.get("applicable_model"),"Universal")), str(clean_value(row.get("reference_image"),"")), to_bool(row.get("requires_oring")), str(clean_value(row.get("oring_spec"),"")), to_bool(row.get("requires_silicone_grease")), str(clean_value(row.get("silicone_grease_note"),"")), to_float(row.get("torque_min")), to_float(row.get("torque_max")), str(clean_value(row.get("notes"),"")))

def validate_components(components):
    errors, warnings, seen = [], [], set()
    for i, c in enumerate(components, 1):
        if not c.part_name: errors.append(f"第 {i} 列缺少 part_name")
        if not c.part_number: errors.append(f"第 {i} 列缺少 part_number")
        if not c.size: errors.append(f"第 {i} 列缺少 size")
        if c.part_number in seen: errors.append(f"重複 part_number：{c.part_number}")
        seen.add(c.part_number)
        target, unit = c.required_torque is not None, bool(c.torque_unit)
        if target != unit: errors.append(f"{c.part_number} 的 required_torque 與 torque_unit 必須同時填寫或同時留空")
        elif not target: warnings.append(f"{c.part_number}：未設定扭力，視為無需扭力元件")
        elif c.torque_min is not None and c.torque_max is not None and not (c.torque_min <= c.required_torque <= c.torque_max): errors.append(f"{c.part_number} 的扭力上下限不合理")
    if errors: raise ValueError("\n".join(errors))
    return warnings

def load_components(path):
    path = Path(path); rows=[]
    if path.suffix.lower()==".xlsx":
        wb=load_workbook(path,data_only=True,read_only=True); sheet=wb["components"] if "components" in wb.sheetnames else wb.active; headers=[c.value for c in next(sheet.iter_rows())]; rows=[dict(zip(headers,v)) for v in sheet.iter_rows(min_row=2,values_only=True) if not all(x is None for x in v)]
    elif path.suffix.lower() in {".yaml",".yml"}: rows=(yaml.safe_load(path.read_text(encoding="utf-8")) or {}).get("components",[])
    else: raise ValueError("只支援 .xlsx、.yaml、.yml")
    components=[component_from_dict(row) for row in rows]; return components, validate_components(components)

def export_pptx(path, image_path, components, annotations, window):
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); left,top=Inches(.5),Inches(.5); pic=slide.shapes.add_picture(image_path,left,top,width=Inches(8.5)); image_width=Inches(8.5); image_height=image_width*pic.height/pic.width
    for a in annotations:
        c=components[a.component_index]; x,y=left+int(image_width*a.x),top+int(image_height*a.y); radius=Inches(.12)
        circle=slide.shapes.add_shape(MSO_SHAPE.OVAL,x-radius,y-radius,radius*2,radius*2); circle.fill.background(); circle.line.color.rgb=RGBColor(220,0,0); circle.line.width=Pt(2.5)
        text_box=slide.shapes.add_textbox(x-radius,y-radius,radius*2,radius*2); tf=text_box.text_frame; tf.clear(); p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; run=p.add_run(); run.text=str(a.sequence); run.font.size=Pt(10); run.font.bold=True; run.font.color.rgb=RGBColor(220,0,0)
        tx,ty=x+Inches(.16),y-Inches(.18); name_box=slide.shapes.add_textbox(tx,ty,Inches(2.8),Inches(.35)); p=name_box.text_frame.paragraphs[0]; run=p.add_run(); run.text=c.part_name; run.font.size=Pt(12); run.font.color.rgb=RGBColor(0,0,0)
        torque=window.torque_display(a)
        if torque:
            tb=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,tx,ty+Inches(.35),Inches(2.0),Inches(.35)); tb.fill.background(); tb.line.color.rgb=RGBColor(210,0,0); tb.line.width=Pt(1.5); p=tb.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; run=p.add_run(); run.text=torque; run.font.size=Pt(11); run.font.color.rgb=RGBColor(0,0,0)
        line=slide.shapes.add_connector(1,tx,ty+Inches(.18),x,y); line.line.color.rgb=RGBColor(220,0,0); line.line.width=Pt(2)
    prs.save(path)

def main():
    app=QApplication(sys.argv); window=MainWindow(); window.show(); sys.exit(app.exec())

if __name__=="__main__": main()
